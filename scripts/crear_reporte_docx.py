import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

def crear_reporte_docx():
    script_dir = os.path.dirname(os.path.abspath(__file__))
    proyecto_root = os.path.dirname(script_dir)
    docx_path = os.path.join(proyecto_root, "Anexo_Tecnico_Clustering_y_Despliegue.docx")
    
    print(f"Iniciando la creación del documento Word en: {docx_path}")
    doc = docx.Document()
    
    # ------------------ ESTILOS DE FUENTE ------------------
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Arial'
    style_normal.font.size = Pt(11)
    style_normal.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    
    # Colores institucionales
    COLOR_PRIMARY = RGBColor(11, 19, 43)    # Azul oscuro
    COLOR_SECONDARY = RGBColor(0, 180, 216) # Cian
    
    def add_custom_heading(text, level, space_before=12, space_after=6):
        heading = doc.add_heading(text, level=level)
        heading.paragraph_format.space_before = Pt(space_before)
        heading.paragraph_format.space_after = Pt(space_after)
        heading.paragraph_format.keep_with_next = True
        
        for run in heading.runs:
            run.font.name = 'Arial'
            if level == 1:
                run.font.size = Pt(18)
                run.font.color.rgb = COLOR_PRIMARY
                run.bold = True
            elif level == 2:
                run.font.size = Pt(14)
                run.font.color.rgb = COLOR_SECONDARY
                run.bold = True
            elif level == 3:
                run.font.size = Pt(12)
                run.font.color.rgb = COLOR_PRIMARY
                run.bold = True
                run.italic = True
        return heading

    # ------------------ 1. PORTADA FORMAL DEL TRABAJO ------------------
    for _ in range(3):
        doc.add_paragraph()
        
    p_institution = doc.add_paragraph()
    p_institution.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_inst = p_institution.add_run("UNIVERSIDAD NACIONAL DE SAN LUIS\nFacultad de Ciencias Físico Matemáticas y Naturales\nDepartamento de Informática")
    run_inst.font.size = Pt(12)
    run_inst.font.color.rgb = COLOR_PRIMARY
    run_inst.bold = True
    
    for _ in range(4):
        doc.add_paragraph()
        
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_title.add_run("TABLERO DE CONTROL DE GESTIÓN Y PRIORIZACIÓN DEL DELITO AUTOMOTOR")
    run_title.font.size = Pt(20)
    run_title.font.color.rgb = COLOR_PRIMARY
    run_title.bold = True
    
    p_subtitle = doc.add_paragraph()
    p_subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = p_subtitle.add_run(
        "Documentación Técnica y Fundamentación de Indicadores de Alerta Federal\n"
        "Espacio Curricular: Fundamentos de Ciencia de Datos (2026)\n"
        "Integrantes: Ezequiel Bernaldez, Ezequiel Nodar"
    )
    run_sub.font.size = Pt(11)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(100, 100, 100)
    
    for _ in range(5):
        doc.add_paragraph()
        
    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_meta = p_meta.add_run(
        "Materia: Fundamentos de Ciencia de Datos\n"
        "Estudiantes: Ezequiel Bernaldez, Ezequiel Nodar\n"
        "Cátedra: Prof. Barrionuevo Mercedes, Prof. Arroyuelo Jorge\n"
        "Año: 2026\n"
        "San Luis, Argentina"
    )
    run_meta.font.size = Pt(11)
    run_meta.font.color.rgb = COLOR_PRIMARY
    
    doc.add_page_break()
    
    # ------------------ 2. INTRODUCCIÓN ------------------
    add_custom_heading("1. Portada e Introducción Institucional", level=1)
    
    p = doc.add_paragraph(
        "El presente reporte documenta el desarrollo y despliegue del Sistema Federal de Monitoreo y Priorización "
        "del Delito Automotor (SFM-DA), estructurado como un entregable formal y estratégico de alto nivel dirigido al "
        "Ministerio de Seguridad de la Nación. Su objetivo primordial es consolidar un panel de control interactivo que "
        "elimine los sesgos tradicionales en la evaluación de la eficiencia policial, permitiendo priorizar las políticas "
        "públicas de seguridad y patrullaje preventivo en las 24 jurisdicciones de la República Argentina."
    )
    p.paragraph_format.space_after = Pt(10)
    
    p = doc.add_paragraph(
        "La interfaz del tablero de control de gestión se inicia con una portada formal e institucional interactiva que presenta "
        "el título, los integrantes del grupo técnico (Ezequiel Bernaldez, Ezequiel Nodar), la materia académica de sustento "
        "(Fundamentos de Ciencia de Datos), el año del estudio (2026) y una breve descripción de la consola que introduce "
        "metodológicamente al auditor en los objetivos operativos del Ministerio."
    )
    p.paragraph_format.space_after = Pt(10)
    
    # ------------------ 3. EL ÍNDICE IPI ------------------
    add_custom_heading("2. Fundamentación del Índice de Prioridad de Intervención (IPI)", level=1)
    
    p = doc.add_paragraph(
        "La evaluación tradicional de la eficiencia en recuperación vehicular basada puramente en tasas porcentuales resulta "
        "altamente engañosa para la toma de decisiones ministeriales. Si una jurisdicción pequeña registra 10 robos y recupera "
        "5, figura con una tasa de recupero del 50%. De igual modo, una gran urbe con 10.000 robos que recupera 5.000 vehículos "
        "figura con la misma tasa del 50%. No obstante, a escala gubernamental, el impacto social, económico y de desprotección "
        "asociado a dejar 5.000 vehículos sustraídos en circulación ilegal es inmensamente más crítico que dejar 5."
    )
    p.paragraph_format.space_after = Pt(10)
    
    p = doc.add_paragraph(
        "Para subsanar este sesgo de escala, se definió el Índice de Prioridad de Intervención (IPI), el cual representa el "
        "número físico neto de vehículos sustraídos que no fueron recuperados en el período analizado:"
    )
    p.paragraph_format.space_after = Pt(6)
    
    p_math1 = doc.add_paragraph()
    p_math1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_m1 = p_math1.add_run("IPI = Robos - Recuperos = Robos * (1 - Tasa de Recupero)")
    run_m1.bold = True
    run_m1.font.size = Pt(12)
    p_math1.paragraph_format.space_after = Pt(10)
    
    add_custom_heading("2.1 Transformación Logarítmica Natural y Estandarización", level=2)
    
    p = doc.add_paragraph(
        "Al unificar las bases de datos de la DNRPA para el rango completo de análisis (2020-2026), que acumula "
        "más de 242.000 registros, la brecha de volumen absoluta entre Buenos Aires (156.808 robos acumulados) y "
        "Catamarca o La Rioja (menos de 90 robos) es tan masiva que anula la capacidad del algoritmo de agrupamiento K-Means. "
        "K-Means, al basarse en distancias euclidianas directas, detectará a Buenos Aires como un outlier masivo aislado en "
        "su propio clúster y agrupará a las 23 provincias restantes en un clúster secundario sin discriminación de perfiles."
    )
    p.paragraph_format.space_after = Pt(10)
    
    p = doc.add_paragraph(
        "Para corregir metodológicamente esto, aplicamos una transformación logarítmica natural sobre el IPI neto (añadiendo "
        "una constante +1 para evitar inconsistencias en valores cero):"
    )
    p.paragraph_format.space_after = Pt(6)
    
    p_math2 = doc.add_paragraph()
    p_math2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_m2 = p_math2.add_run("IPI_log = ln(IPI + 1)")
    run_m2.bold = True
    run_m2.font.size = Pt(12)
    p_math2.paragraph_format.space_after = Pt(10)
    
    p = doc.add_paragraph(
        "Posteriormente, aplicamos una estandarización estricta Z-Score sobre IPI_log y la tasa de recupero tradicional. "
        "Esta normalización comprime la dispersión de escalas euclidianas permitiendo agrupar eficientemente provincias "
        "grandes, medianas y pequeñas. En el tablero, el logaritmo es re-escalado linealmente de 0 a 100 para dotar al "
        "Ministerio de un semáforo visual de prioridad federal directa."
    )
    p.paragraph_format.space_after = Pt(10)
    
    # ------------------ 4. ARQUITECTURA FRONTE-END ------------------
    add_custom_heading("3. Estructura y Rediseño de la Consola de Mando", level=1)
    
    p = doc.add_paragraph(
        "La interfaz del tablero de control de gestión fue reestructurada para ajustarse a los requerimientos formales de "
        "evaluación de organismos gubernamentales, optimizando la visibilidad de datos y eliminando ruidos en la navegación:"
    )
    p.paragraph_format.space_after = Pt(10)
    
    p = doc.add_paragraph(
        "• Remoción del Módulo CRISP-DM: Se retiró el flujo metodológico interactivo para centrar la consola de mando en "
        "la visualización analítica pura de gestión, trasladando las justificaciones técnicas a este informe.\n"
        "• Eliminación de la Ficha Lateral: Se removió la ficha de control provincial fija que ocupaba espacio de visualización lateral "
        "en la Consola de Mando.\n"
        "• Extensión Horizontal de la Tabla de Prioridades (100% Ancho): La tabla 'Ranking Federal de Alerta y Prioridad' "
        "fue extendida horizontalmente a todo lo ancho de la pantalla, facilitando al auditor la lectura lineal completa de robos, "
        "recuperos, tasas de recupero reales, IPIs acumulados, porcentaje de alerta (0-100) y categoría de cluster de cada provincia."
    )
    p.paragraph_format.space_after = Pt(10)
    
    # ------------------ 5. SERIE TEMPORAL POR PROVINCIA ------------------
    add_custom_heading("4. Módulo Explorador de Series Temporales Provinciales", level=1)
    
    p = doc.add_paragraph(
        "Para responder a la necesidad de evaluar el comportamiento histórico local en cada jurisdicción, incorporamos en la pestaña "
        "de Análisis de Tendencias un 'Explorador Temporal Provincial' interactivo. Este componente consta de un menú de búsqueda "
        "desplegable conteniendo las 24 provincias normalizadas y un gráfico dinámico multi-eje en Chart.js."
    )
    p.paragraph_format.space_after = Pt(10)
    
    p = doc.add_paragraph(
        "Al seleccionar cualquier provincia, el sistema recupera su historial mensual agregando sus variables a lo largo del tiempo "
        "(2020-2026) y grafica:\n"
        "1. Denuncias de Robo (Eje vertical izquierdo, curva roja con sombreado de relleno).\n"
        "2. Comunicaciones de Recupero (Eje vertical derecho, curva verde agua sólida).\n"
        "3. Evolución de la Tasa de Recuperación porcentual a nivel mensual (Curva de trazos amarillos discretos).\n"
        "Esto dota al Ministerio de un análisis de tendencias locales 100% interactivo y de alta granularidad temporal."
    )
    p.paragraph_format.space_after = Pt(10)
    
    # ------------------ 6. CONCENTRACIÓN CON PORCENTAJES ------------------
    add_custom_heading("5. Reporte de Concentración Delictiva con Porcentajes Exactos", level=1)
    
    p = doc.add_paragraph(
        "El análisis exploratorio en la pestaña de tendencias reporta ahora el porcentaje exacto de incidencia delictiva "
        "que cada una de las 24 provincias representa sobre el volumen total nacional acumulado. Este reporte combina un gráfico "
        "de rosca de alta fidelidad visual y una tabla interactiva con barras de progreso de micro-escala:"
    )
    p.paragraph_format.space_after = Pt(10)
    
    # Tabla representativa de concentraciones
    table_pct = doc.add_table(rows=6, cols=3)
    table_pct.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_pct.autofit = True
    
    hdr_cells = table_pct.rows[0].cells
    hdr_cells[0].text = 'Jurisdicción'
    hdr_cells[1].text = 'Robos Acumulados'
    hdr_cells[2].text = 'Porcentaje de Incidencia Nacional (%)'
    
    for cell in hdr_cells:
        cell.paragraphs[0].runs[0].font.bold = True
        cell.paragraphs[0].runs[0].font.size = Pt(10)
        cell.paragraphs[0].runs[0].font.color.rgb = COLOR_PRIMARY
        
    rows_pct = [
        ('BUENOS AIRES', '156.808', '67,02%'),
        ('CIUDAD AUTÓNOMA DE BUENOS AIRES', '28.434', '12,15%'),
        ('CÓRDOBA', '19.631', '8,39%'),
        ('SANTA FE', '11.253', '4,81%'),
        ('MENDOZA', '10.977', '4,69%')
    ]
    
    for i, data_row in enumerate(rows_pct):
        cells = table_pct.rows[i+1].cells
        cells[0].text = data_row[0]
        cells[1].text = data_row[1]
        cells[2].text = data_row[2]
        
    for row in table_pct.rows[1:]:
        for cell in row.cells:
            cell.paragraphs[0].runs[0].font.size = Pt(9.5)
            
    doc.add_paragraph().paragraph_format.space_after = Pt(10)
    
    p = doc.add_paragraph(
        "Este desglose porcentual exhaustivo evidencia una hiper-concentración del delito: solo 5 provincias concentran el 97.06% "
        "de la sustracción del parque automotor de toda la República Argentina, justificando de forma inmediata la centralización "
        "geográfica de presupuestos especiales de seguridad pública."
    )
    p.paragraph_format.space_after = Pt(10)
    
    # ------------------ 7. MODELO CLUSTERING ------------------
    add_custom_heading("6. Modelo de Machine Learning K-Means y Clasificación de Riesgo", level=1)
    
    p = doc.add_paragraph(
        "Al entrenar el modelo K-Means con los datos consolidados y normalizados de las 24 provincias de la República Argentina, "
        "el coeficiente de silueta promedio alcanzó un valor estadísticamente excelente de 0.5625, lo que valida la rigurosidad "
        "científica del agrupamiento. Las provincias se clasificaron en tres dinámicas estables:"
    )
    p.paragraph_format.space_after = Pt(10)
    
    p = doc.add_paragraph(
        "1. Clúster de Nivel de Intervención Crítico (5 Provincias): Buenos Aires, CABA, Córdoba, Mendoza y Santa Fe. "
        "Se caracterizan por prioridades máximas en el tablero (Buenos Aires en 100/100, CABA en 76.5/100, Córdoba en 72.5/100), "
        "masivos volúmenes y tasas de recupero bajas en relación a su flujo. Requieren urgentes comités de seguridad integrados y control "
        "de venta de autopartes online.\n"
        "2. Clúster de Nivel de Intervención Moderado (1 Provincia): Entre Ríos. Presenta un volumen delictivo intermedio (698 robos) "
        "pero destaca a escala federal con la tasa de recuperación más eficiente de la República Argentina (18.91%), lo que la ubica en "
        "un clúster particular de alta efectividad.\n"
        "3. Clúster de Nivel de Eficiencia Destacada / Controlado (18 Provincias): San Luis, Neuquén, Chubut, Salta, Jujuy, Chaco, etc. "
        "Jurisdicciones con muy baja concentración delictiva y riesgos controlados. San Luis obtiene una prioridad de 19.0/100 en la "
        "consola, plenamente controlada bajo la estructura de seguridad ordinaria."
    )
    p.paragraph_format.space_after = Pt(10)
    
    # ------------------ 8. CONCLUSIONES ------------------
    add_custom_heading("7. Conclusiones Estratégicas para Seguridad Pública", level=1)
    
    p = doc.add_paragraph(
        "• Redirección de Patrullaje: El IPI logarítmico consolidó un mapa temático que guía al Ministerio para desplegar operativos de "
        "saturación en el corredor Buenos Aires-CABA de forma prioritaria, optimizando recursos fiscales.\n"
        "• Replicabilidad del Caso Entre Ríos: El modelo identificó a Entre Ríos en un perfil destacado de recuperación vehicular. Estudiar "
        "su estrategia de cerrojo fronterizo vial e interprovincial permitirá modelar políticas aplicables a Santa Fe o Córdoba.\n"
        "• Solidez de Datos: La normalización de las 24 provincias y el capping de outliers cuantitativos en Python eliminaron duplicaciones "
        "engañosas en CABA, dotando de trazabilidad absoluta y reproducibilidad científica al estudio federal interanual."
    )
    p.paragraph_format.space_after = Pt(10)
    
    doc.save(docx_path)
    print(f"¡Documento Word Anexo_Tecnico_Clustering_y_Despliegue.docx creado y guardado exitosamente!")

if __name__ == "__main__":
    crear_reporte_docx()
